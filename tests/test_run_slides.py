"""slides_export step 3 (ADR-165): one click from experiment results - a run's image artifacts appear as figures
in the run deck exported through the host (`POST /api/v1/exports/report_slides`)."""

from __future__ import annotations

import base64
import hashlib
import sqlite3
from pathlib import Path
from typing import Any

import pytest

from researchos.kernel.config import Settings
from researchos.kernel.db import new_id, utcnow
from researchos.services.container import Services
from researchos.services.execution.ingest import _kind

from .conftest import Client
from .fixtures.make_dataset import make
from .m2_helpers import new_plan, plan_content, run

PIXEL = Path(__file__).parents[1] / "plugins" / "report_export" / "fixtures" / "pixel.b64"


def test_run_deck_shows_run_figures(client: Client, conn: sqlite3.Connection, svc: Services, settings: Settings,
                                    tmp_path: Path) -> None:
    pytest.importorskip("researchos_plugin_report_export")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from researchos_plugin_report_export.slides import SlidesExporter

    assert _kind("figures/accuracy.png") == "figure" and _kind("figures/x.JPG") == "figure"
    ds = client.post("/api/v1/datasets", json={"path": str(make(tmp_path / "syn"))}).json()
    t = run(client, conn, svc, new_plan(client, plan_content(client, ds)))
    r: dict[str, Any] = t["run"]
    png = base64.b64decode(PIXEL.read_text())
    figure = settings.data_dir / r["run_dir"] / "figures" / "accuracy.png"
    figure.parent.mkdir(parents=True)
    figure.write_bytes(png)  # as a template would; ingest registers files with this kind (asserted above)
    conn.execute("INSERT INTO artifacts(id,run_id,path,sha256,size,kind,created_at) VALUES (?,?,?,?,?,?,?)",
                 (new_id("art"), r["id"], figure.relative_to(settings.data_dir).as_posix(),
                  hashlib.sha256(png).hexdigest(), len(png), "figure", utcnow()))
    project = client.post("/api/v1/projects", json={"title": "P"}).json()
    conn.execute("UPDATE plans SET project_id=?", (project["id"],))
    svc.extensions.exporters["report_slides"] = SlidesExporter()
    out = client.post("/api/v1/exports/report_slides", json={"run_id": r["id"], "mode": "final"}).json()
    assert out["status"] == "ok", out
    ev = conn.execute("SELECT project_id FROM events WHERE kind='export.created' AND entity_id=?", (out["id"],))
    assert ev.fetchone()["project_id"] == project["id"]  # review #56: exports record the run's project
    deck = Presentation(str(settings.data_dir / next(f["path"] for f in out["files"] if f["path"].endswith(".pptx"))))
    pictures = [s for slide in deck.slides for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    texts = " ".join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert len(pictures) == 3 and png in [p.image.blob for p in pictures]  # + 2 host figures (ADR-194)
    assert "accuracy.png (run artifact)" in texts and "confusion_search_val.png" in texts and any(
        s.has_table for slide in deck.slides for s in slide.shapes)  # the metrics table is native


def test_export_download_serves_only_registered_unchanged_files(client: Client, conn: sqlite3.Connection,
                                                               settings: Settings, tmp_path: Path) -> None:
    """ADR-167/172 (review #56): download by export id + file name; only files of a registered, non-failed export
    whose bytes still match the recorded hash; never arbitrary paths under exports/."""
    from researchos.kernel.events import emit

    d = settings.data_dir / "exports" / "doc-report_slides-x"
    d.mkdir(parents=True)
    deck = d / "report.pptx"
    deck.write_bytes(b"PK deck")
    (d / "stray.pptx").write_bytes(b"PK unregistered")
    rel = deck.relative_to(settings.data_dir).as_posix()
    files = [{"path": rel, "mime": "pptx", "sha256": hashlib.sha256(b"PK deck").hexdigest()}]
    emit(conn, "export.created", entity_type="export", entity_id="exp_ok", data={"status": "ok", "files": files})
    emit(conn, "export.created", entity_type="export", entity_id="exp_bad", data={"status": "failed", "files": files})
    ok = client.get("/api/v1/exports/exp_ok/files/report.pptx")
    assert ok.status_code == 200 and ok.content == b"PK deck" and "attachment" in ok.headers["content-disposition"]
    for url in ("/api/v1/exports/exp_ok/files/stray.pptx", "/api/v1/exports/exp_bad/files/report.pptx",
                "/api/v1/exports/nope/files/report.pptx", "/api/v1/exports/download?path=" + rel):
        assert client.get(url).status_code == 404, url
    deck.write_bytes(b"PK replaced")  # bytes no longer match the record
    assert client.get("/api/v1/exports/exp_ok/files/report.pptx").status_code == 404

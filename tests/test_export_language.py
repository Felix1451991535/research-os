"""Export language (ADR-234): the run report exported as Markdown and as .pptx in en and zh-CN from the SAME run has
identical numbers (verification labels included) and passes the number gate in both; the zh export carries no English
template text from the catalogs; an unsupported language is refused. Catalog parity / placeholder checks."""

from __future__ import annotations

import re
import sqlite3
from pathlib import Path
from typing import Any

import pytest

from researchos.kernel.config import Settings
from researchos.sdk.texts import Catalog, problems
from researchos.services.container import Services
from researchos.services.content.texts import CORE

from .conftest import Client
from .fixtures.make_dataset import make
from .m2_helpers import new_plan, plan_content, run

NUMBER = re.compile(r"\d+\.\d{4}(?: \w+)? \((?:recomputed|invariant_checked|self_reported|[a-z_]+)\)")


@pytest.fixture
def run_id(client: Client, conn: sqlite3.Connection, svc: Services, tmp_path: Path) -> str:
    ds = client.post("/api/v1/datasets", json={"path": str(make(tmp_path / "syn"))}).json()
    return str(run(client, conn, svc, new_plan(client, plan_content(client, ds)))["run"]["id"])


def _export(client: Client, exporter: str, run_id: str, language: str | None) -> dict[str, Any]:
    options = {"language": language} if language else {}
    r = client.post(f"/api/v1/exports/{exporter}", json={"run_id": run_id, "mode": "final", "options": options})
    assert r.status_code == 200, r.text
    out: dict[str, Any] = r.json()
    assert out["status"] == "ok" and not [d for d in out["diagnostics"] if d["severity"] == "error"], out
    return out


def _file(settings: Settings, out: dict[str, Any], suffix: str) -> Path:
    return settings.data_dir / next(f["path"] for f in out["files"] if f["path"].endswith(suffix))


def test_markdown_run_report_in_both_languages_has_identical_numbers(client: Client, settings: Settings,
                                                                      run_id: str) -> None:
    en = _file(settings, _export(client, "markdown", run_id, None), ".md").read_text(encoding="utf-8")
    zh = _file(settings, _export(client, "markdown", run_id, "zh-CN"), ".md").read_text(encoding="utf-8")
    assert "## Metrics" in en and "## 指标" in zh and f"# 运行报告 {run_id}" in zh
    assert NUMBER.findall(en) and NUMBER.findall(en) == NUMBER.findall(zh)
    english = [v for k, v in CORE.texts["en"].items() if "{" not in v and v not in CORE.texts["zh-CN"].values()]
    leaked = [v for v in english if re.search(rf"(?m)^#+ {re.escape(v)}$|\| {re.escape(v)} \|", zh)]
    assert not leaked, leaked


def test_slides_run_report_in_both_languages(client: Client, svc: Services, settings: Settings, run_id: str) -> None:
    pytest.importorskip("researchos_plugin_report_export")
    from pptx import Presentation
    from researchos_plugin_report_export.slides import SlidesExporter

    svc.extensions.exporters["report_slides"] = SlidesExporter()

    def deck(language: str) -> tuple[str, list[str]]:
        d = Presentation(str(_file(settings, _export(client, "report_slides", run_id, language), ".pptx")))
        texts = " ".join(s.text for slide in d.slides for s in slide.shapes if s.has_text_frame)
        cells = [c.text for slide in d.slides for s in slide.shapes if s.has_table for r in s.table.rows
                 for c in r.cells]
        return texts, cells

    en_text, en_cells = deck("en")
    zh_text, zh_cells = deck("zh-CN")
    assert "Report" in en_text and "报告" in zh_text and "数字和引用均链接到已核验的记录" in zh_text
    assert "背景" in zh_text and "Background" not in zh_text and "Numbers and citations" not in zh_text
    numbers = [c for c in en_cells if NUMBER.fullmatch(c)]
    assert numbers and numbers == [c for c in zh_cells if NUMBER.fullmatch(c)]
    assert en_cells[:3] == ["metric", "split", "value"] and zh_cells[:3] == ["指标", "数据划分", "数值"]


def test_unsupported_language_is_refused(client: Client, run_id: str) -> None:
    r = client.post("/api/v1/exports/markdown", json={"run_id": run_id, "options": {"language": "fr"}})
    assert r.status_code == 422 and r.json()["code"] == "bad_language"


def test_catalog_checks_parity_empty_and_placeholders() -> None:
    assert problems(CORE.texts) == []
    bad = {"en": {"a": "x {n}", "b": "y", "c": "z"}, "zh-CN": {"a": "x {m}", "b": " ", "d": "w"}}
    assert problems(bad) == ["zh-CN: missing c", "zh-CN: extra d", "zh-CN: empty b", "zh-CN: placeholders of a differ"]
    with pytest.raises(ValueError, match="invalid text catalog"):
        Catalog(bad)
    assert problems({"en": {}}) == ["missing language zh-CN"]

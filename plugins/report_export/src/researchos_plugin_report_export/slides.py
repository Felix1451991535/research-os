"""Slides exporter `report_slides`: readable research slides with an evidence provenance sidecar (ADR-132, ADR-160).
Rendering lives here; resolution (citations, metrics, artifacts, diagnostics) is shared with the other exporters."""

from __future__ import annotations

import io
import re
import unicodedata
from collections.abc import Callable, Mapping
from datetime import UTC, datetime
from typing import Any, cast

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.presentation import Presentation as PptxPresentation
from pptx.slide import Slide
from pptx.util import Emu
from pptx.util import Inches as SlideInches
from pptx.util import Pt as SlidePt
from researchos.contracts.exporters import (
    Diagnostic,
    ExportContext,
    ExporterSpec,
    ExportFile,
    ExportResult,
    ReportSnapshot,
)

from .exporters import provenance, validate_options
from .resolve import Resolved, clean, references, resolve, safe_output, slide_sections

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
RGB = cast(Callable[[int, int, int], RGBColor], RGBColor)
SLIDE_OPTIONS: dict[str, Any] = {"type": "object", "properties": {}, "additionalProperties": False}


# Line breaking by glyph advance (em) instead of character count (review #14): printable-ASCII advances of
# Arial (measured on metric-compatible Liberation Sans, rounded up); East Asian wide = 1 em; any other character is
# taken as a conservative 0.95 em. 24 pt body text in the 12.3 in box (0.1 in insets each side) -> 36.3 em; budget
# 35.5 em keeps a margin.
_ADVANCE = {
    0.191: "'",
    0.223: "ijl",
    0.26: "|",
    0.278: " !,./:;I[\\]ft",
    0.334: "()-`r{}",
    0.355: '"',
    0.39: "*",
    0.47: "^",
    0.5: "Jcksvxyz",
    0.557: "#$0123456789?L_abdeghnopqu",
    0.584: "+<=>~",
    0.611: "FTZ",
    0.667: "&ABEKPSVXY",
    0.723: "CDHNRUw",
    0.778: "GOQ",
    0.834: "Mm",
    0.89: "%",
    0.944: "W",
    1.016: "@",
}
EM = {c: w for w, chars in _ADVANCE.items() for c in chars}
LINE_EM = 35.5


def _width(text: str) -> float:
    return sum(
        EM.get(c, 0.0 if unicodedata.combining(c) else (1.0 if unicodedata.east_asian_width(c) in "WF" else 0.95))
        for c in text
    )


def _hard_split(token: str, width: float) -> list[str]:
    out, line = [], ""
    for char in token:
        if _width(line + char) > width:
            out.append(line)
            line = ""
        line += char
    return [*out, line]


def slide_lines(value: str, width: float = LINE_EM) -> list[str]:
    """Wrap at word boundaries within `width`; only a token wider than a line (e.g. unspaced CJK) is split.
    Text is never shrunk."""
    lines: list[str] = []
    for source in clean(value).expandtabs(4).splitlines():
        line = ""
        for word in source.split(" "):
            candidate = f"{line} {word}" if line else word
            if _width(candidate) <= width:
                line = candidate
                continue
            if line:
                lines.append(line)
            *full, line = _hard_split(word, width)
            lines.extend(full)
        lines.append(line)
    return lines


def slide_text(slide: Slide, lines: list[str], *, heading: bool = False) -> None:
    box = slide.shapes.add_textbox(
        SlideInches(0.5), SlideInches(0.3 if heading else 1.2), SlideInches(12.3), SlideInches(0.7 if heading else 5.8)
    )
    _fill(box.text_frame, lines, 28 if heading else 24)


def _fill(frame: Any, lines: list[str], size: int, color: tuple[int, int, int] = (32, 32, 32)) -> None:
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.space_before = SlidePt(0)
        paragraph.space_after = SlidePt(0)
        paragraph.line_spacing = SlidePt(size * 1.5)
        run = paragraph.add_run()
        run.text = line
        run.font.size = SlidePt(size)
        run.font.name = "Arial"
        run.font.color.rgb = RGB(*color)


def slides_page(deck: PptxPresentation, title: str) -> Slide:
    slide = cast(Slide, deck.slides.add_slide(deck.slide_layouts[6]))
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGB(255, 255, 255)
    slide_text(slide, [title], heading=True)
    return slide


BODY_LINES = 9  # + the title line = 10 lines per slide


def _pages(groups: list[list[str]]) -> list[list[str]]:
    """Fill slides group by group (a list item or paragraph never straddles two slides unless it alone is taller
    than a slide)."""
    pages: list[list[str]] = [[]]
    for group in (g for g in groups if g):
        if pages[-1] and len(pages[-1]) + len(group) > BODY_LINES:
            pages.append([])
        for line in group:
            if len(pages[-1]) == BODY_LINES:
                pages.append([])
            pages[-1].append(line)
    return [p for p in pages if p]


def slides_text_pages(deck: PptxPresentation, title: str, groups: list[list[str]], result: Resolved) -> None:
    pages = _pages(groups)
    if len(pages) > 1:
        result.diagnostics.append(
            Diagnostic("warning", "slide_overflow", f"{title}: content split across {len(pages)} slides", title)
        )
    for i, page in enumerate(pages):
        slide_text(slides_page(deck, title + (" (continued)" if i else "")), page)


TABLE_PT = 16  # cell text size
LINE_PT, PAD_PT = TABLE_PT * 1.25, 8.0  # one wrapped line; top + bottom cell margins
TABLE_TOP_IN, TABLE_BOTTOM_IN = 1.2, 6.9  # content box: below the title, above the footer
BUDGET_PT = (TABLE_BOTTOM_IN - TABLE_TOP_IN) * 72
NUMBER = re.compile(r"^[-+\u2212]?(\d[\d,]*)?(\.\d+)?([eE][-+]?\d+)?\s*%?$")
Cells = list[list[str]]  # one row: wrapped lines per cell


def _cell_lines(text: str, columns: int) -> list[str]:
    return slide_lines(text, LINE_EM * 24 / TABLE_PT / columns - 1.0) or [""]


def _height(row: Cells) -> float:
    return max(len(c) for c in row) * LINE_PT + PAD_PT


def _fragments(row: Cells, max_lines: int) -> list[Cells]:
    """A row taller than a slide continues on the next slide in line-aligned pieces (review #62)."""
    n = max(len(c) for c in row)
    return [[c[k : k + max_lines] or [""] for c in row] for k in range(0, n, max_lines)]


def _table_pages(rows: list[list[str]]) -> list[list[Cells]]:
    """Header + rows per slide within the content box, by each row's own wrapped height; a row splits only when it
    alone is taller than the space under the header."""
    cols = len(rows[0])
    header, *body = [[_cell_lines(clean(c), cols) for c in row] for row in rows]
    room = BUDGET_PT - _height(header)
    max_lines = max(1, int((room - PAD_PT) // LINE_PT))
    pages: list[list[Cells]] = [[]]
    used = 0.0
    for row in (f for r in body for f in _fragments(r, max_lines)):
        if pages[-1] and used + _height(row) > room:
            pages.append([])
            used = 0.0
        pages[-1].append(row)
        used += _height(row)
    return [[header, *page] for page in pages]


def table_slides(deck: PptxPresentation, title: str, rows: list[list[str]], result: Resolved) -> None:
    """Native table (header row bold, numbers right-aligned), paged by rows, each row sized to its wrapped text so
    the table ends above the footer; PowerPoint/WPS-portable plain table."""
    pages = _table_pages(rows)
    if len(pages) > 1:
        result.diagnostics.append(
            Diagnostic("warning", "slide_overflow", f"{title}: table split over {len(pages)} slides", title)
        )
    for i, page in enumerate(pages):
        slide = slides_page(deck, title + (" (continued)" if i else ""))
        heights = [_height(row) for row in page]
        shape = slide.shapes.add_table(
            len(page),
            len(page[0]),
            SlideInches(0.5),
            SlideInches(TABLE_TOP_IN),
            SlideInches(12.3),
            Emu(int(SlidePt(sum(heights)))),
        )
        _plain_grid(shape)
        for r, (row, h) in enumerate(zip(page, heights, strict=True)):
            shape.table.rows[r].height = Emu(int(SlidePt(h)))
            for c, lines in enumerate(row):
                cell = shape.table.cell(r, c)
                cell.text = "\n".join(lines)
                numeric = r > 0 and len(lines) == 1 and bool(NUMBER.match(lines[0].strip())) and lines[0].strip() != ""
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment = PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT
                    for run in paragraph.runs:
                        run.font.size, run.font.name, run.font.bold = SlidePt(TABLE_PT), "Arial", r == 0


PLAIN_GRID = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # built-in "No Style, Table Grid" (PowerPoint and WPS)


def _plain_grid(shape: Any) -> None:
    """Thin black grid, no banding or fill: the academic table look, using a built-in style id only."""
    tbl = shape._element.graphic.graphicData.tbl
    tbl.tblPr.set("firstRow", "1")
    tbl.tblPr.set("bandRow", "0")
    style = tbl.tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
    if style is None:
        from pptx.oxml.ns import qn

        style = tbl.tblPr.makeelement(qn("a:tableStyleId"), {})
        tbl.tblPr.append(style)
    style.text = PLAIN_GRID


def figure_slide(deck: PptxPresentation, title: str, image: bytes, caption: str) -> list[str]:
    """Picture centred in the content box with up to two caption lines under it; returns the caption overflow."""
    slide = slides_page(deck, title)
    picture = slide.shapes.add_picture(io.BytesIO(image), Emu(0), Emu(0))
    scale = min(SlideInches(12.3) / picture.width, SlideInches(4.6) / picture.height)
    picture.width = Emu(int(picture.width * scale))
    picture.height = Emu(int(picture.height * scale))
    picture.left = Emu(int((deck.slide_width - picture.width) / 2))
    picture.top = SlideInches(1.2)
    lines = slide_lines(caption) if caption else []
    if lines:
        box = slide.shapes.add_textbox(SlideInches(0.5), SlideInches(5.95), SlideInches(12.3), SlideInches(0.9))
        box.name = "Caption"
        _fill(box.text_frame, lines[:2], 24)
    return lines[2:]


def _block_group(block: dict[str, Any]) -> list[str]:
    lines: list[str] = []
    if block["kind"] == "figure" and not block.get("image"):
        lines.append("[unresolved figure]")
    if block["kind"] == "list":
        body = slide_lines(str(block["text"]), LINE_EM - _width("\u2022 "))
        lines += ["\u2022 " + body[0], *("  " + x for x in body[1:])] if body and body[0] else []
    else:
        lines += slide_lines(str(block["text"]))
    if block["claim_id"]:
        lines += slide_lines("Claim: " + str(block["claim_id"]))
    if not block.get("image"):  # a figure's caption sits under its picture
        lines += slide_lines(str(block["caption"])) if block["caption"] else []
    return [x for x in lines if x.strip()]


def slides_section(deck: PptxPresentation, title: str, blocks: list[dict[str, Any]], result: Resolved) -> None:
    initial_slides = len(deck.slides)
    groups: list[list[str]] = []
    for block in blocks:
        if block.get("image") or (block["kind"] == "table" and block["rows"]):
            slides_text_pages(deck, title, groups, result)
            groups = []
        if block.get("image"):
            groups.append(figure_slide(deck, title, result.images[block["image"]], str(block["caption"] or "")))
        elif block["kind"] == "table" and block["rows"]:
            table_slides(deck, title, [list(r) for r in block["rows"]], result)
        groups.append(_block_group(block))
    slides_text_pages(deck, title, groups, result)
    if len(deck.slides) == initial_slides:  # an empty section gets no slide; the diagnostics still say so
        result.diagnostics.append(Diagnostic("info", "empty_section", f"{title}: no content", title))


KIND_LABEL = {"report": "Report", "answer": "Cited answer", "manuscript": "Manuscript", "paper_card": "Paper card"}


def title_slide(deck: PptxPresentation, snapshot: ReportSnapshot, result: Resolved) -> None:
    lines = slide_lines(snapshot.title) or ["Untitled"]
    lines += [
        "",
        datetime.now(UTC).strftime("%d %B %Y"),
        "Numbers and citations link to verified records (provenance.json)",
    ]
    slides_text_pages(deck, KIND_LABEL.get(snapshot.kind, "Report"), [lines], result)


def number_slides(deck: PptxPresentation, short_title: str) -> None:
    """Footer on every slide after the title slide: short title and "n / N" (12 pt, grey; named "Footer")."""
    total = len(deck.slides)
    label = (slide_lines(short_title, 24.0) or [""])[0]
    for n, slide in enumerate(deck.slides, start=1):
        if n == 1:
            continue
        box = slide.shapes.add_textbox(SlideInches(0.5), SlideInches(7.0), SlideInches(12.3), SlideInches(0.35))
        box.name = "Footer"
        _fill(box.text_frame, [f"{label}    {n} / {total}"], 12, (110, 110, 110))
        box.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


class SlidesExporter:
    spec = ExporterSpec(
        "report_slides",
        "0.1.0",
        ("report", "answer", "manuscript", "paper_card"),
        (PPTX, "application/json"),
        "Readable research slides with evidence provenance sidecar",
        SLIDE_OPTIONS,
    )

    def export(self, snapshot: ReportSnapshot, ctx: ExportContext, options: Mapping[str, Any]) -> ExportResult:
        validate_options(options, SLIDE_OPTIONS)
        result = resolve(snapshot, ctx)
        if any(d.severity == "error" for d in result.diagnostics):
            return ExportResult((), tuple(result.diagnostics))
        deck = Presentation()
        deck.slide_width, deck.slide_height = SlideInches(13.333), SlideInches(7.5)
        title_slide(deck, snapshot, result)
        for title, blocks in slide_sections(result).items():
            slides_section(deck, title, blocks, result)
        slides_text_pages(deck, "References", [slide_lines(r) for r in references(result)], result)
        number_slides(deck, snapshot.title)
        output = io.BytesIO()
        deck.save(output)
        safe_output(ctx.out_dir, "report.pptx").write_bytes(output.getvalue())
        return ExportResult(
            (ExportFile("report.pptx", PPTX), provenance(snapshot, result, ctx)), tuple(result.diagnostics)
        )

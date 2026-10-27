"""Evidence preservation, final-mode failure, escaping, figures and template layouts."""

import hashlib
import json
from dataclasses import replace
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from researchos.contracts.exporters import ArtifactRef, Block, CiteSpan, MetricSpan, MetricValue, TextSpan
from test_report_export_conformance import context, snapshot

from researchos_plugin_report_export.exporters import DocxExporter, LatexExporter
from researchos_plugin_report_export.latex import escape
from researchos_plugin_report_export.slides import SlidesExporter


@pytest.mark.parametrize("cls", [DocxExporter, LatexExporter])
@pytest.mark.parametrize("template", ["lab_report", "literature_review", "short_paper"])
def test_templates_and_evidence(cls: Any, template: str, tmp_path: Path) -> None:
    result = cls().export(snapshot(), context(tmp_path), {"template": template})
    assert result.files and not result.diagnostics
    provenance = json.loads((tmp_path / "provenance.json").read_text())
    assert provenance["citations"]["test2020"]["evidence"][0]["evidence_id"] == "evidence-1"
    assert provenance["snapshot"]["blocks"][1]["claim_id"] == "claim-1"


@pytest.mark.parametrize("cls", [DocxExporter, LatexExporter, SlidesExporter])
def test_missing_metric_in_table_fails_without_output(cls: Any, tmp_path: Path) -> None:
    report = replace(snapshot(), blocks=(Block("table", rows=(((MetricSpan("bad", "metric", "test"),),),)),))
    result = cls().export(report, context(tmp_path), {})
    assert not result.files and any(d.code == "unresolved_metric" for d in result.diagnostics)
    assert not list(tmp_path.iterdir())


def test_nonfinite_metric(tmp_path: Path) -> None:
    ctx = context(tmp_path)
    ctx.metrics = {("run", "accuracy", "search_val"): MetricValue(float("nan"), None, "self_reported")}
    assert not LatexExporter().export(snapshot(), ctx, {}).files


def test_draft_warning(tmp_path: Path) -> None:
    ctx = context(tmp_path)
    ctx.mode = "draft"
    report = replace(snapshot(), blocks=(Block("paragraph", (CiteSpan("missing"),)),))
    result = LatexExporter().export(report, ctx, {})
    assert result.files and result.diagnostics[0].severity == "warning"


def test_latex_injection_escaped(tmp_path: Path) -> None:
    payload = r"\input{/etc/passwd} % $ & # _ ~ ^"
    report = replace(snapshot(), title=payload)
    LatexExporter().export(report, context(tmp_path), {})
    output = (tmp_path / "report.tex").read_text()
    assert payload not in output and escape(payload) in output


def test_docx_font_xml_and_unicode(tmp_path: Path) -> None:
    title = "\u4eba\u4f53\u6d3b\u52a8\u8bc6\u522b Synthetic"
    DocxExporter().export(replace(snapshot(), title=title), context(tmp_path), {})
    assert Document(str(tmp_path / "report.docx")).paragraphs[0].text == title
    with ZipFile(tmp_path / "report.docx") as z:
        assert b'w:eastAsia="SimSun"' in z.read("word/styles.xml")


def test_artifact_mismatch(tmp_path: Path) -> None:
    image = tmp_path / "source.png"
    image.write_bytes(b"synthetic invalid image")
    ctx = context(tmp_path)
    ctx.artifacts = {"figure": ArtifactRef("figure", "image/png", "0" * 64, "run", image)}
    report = replace(snapshot(), blocks=(Block("figure", artifact_id="figure"),))
    result = LatexExporter().export(report, ctx, {})
    assert not result.files and result.diagnostics[0].code == "artifact_hash_mismatch"


def test_pdf_gap_reported(tmp_path: Path) -> None:
    result = LatexExporter().export(snapshot(), context(tmp_path), {"pdf": True})
    assert result.files and any(d.code == "pdf_unavailable" for d in result.diagnostics)


def test_symlink_output_rejected(tmp_path: Path) -> None:
    outside = tmp_path.parent / "must-not-write.tex"
    (tmp_path / "report.tex").symlink_to(outside)
    with pytest.raises(ValueError):
        LatexExporter().export(snapshot(), context(tmp_path), {})
    assert not outside.exists()


def test_figure_hash_and_caption(tmp_path: Path) -> None:
    import base64

    image = Path(__file__).parents[1] / "fixtures" / "pixel.b64"
    raw = base64.b64decode(image.read_text())
    source = tmp_path / "source.png"
    source.write_bytes(raw)
    ctx = context(tmp_path)
    ctx.artifacts = {"figure": ArtifactRef("figure", "image/png", hashlib.sha256(raw).hexdigest(), "run", source)}
    report = replace(
        snapshot(), blocks=(Block("figure", artifact_id="figure", caption=(TextSpan("Synthetic pixel"),)),)
    )
    result = LatexExporter().export(report, ctx, {})
    assert not result.diagnostics and any(f.mime == "image/png" for f in result.files)


def test_repeated_work_single_reference_all_evidence_preserved(tmp_path: Path) -> None:
    from researchos.contracts.exporters import CitationEntry, EvidenceRef

    ctx = context(tmp_path)
    original = ctx.citations["test2020"]
    alternate = CitationEntry(
        "alternate", original.record, (EvidenceRef("evidence-2", "version-1", 3, "Other quote", "untrusted"),)
    )
    ctx.citations = {**ctx.citations, "alternate": alternate}
    report = replace(snapshot(), blocks=(Block("paragraph", (CiteSpan("test2020"), CiteSpan("alternate"))),))
    LatexExporter().export(report, ctx, {})
    assert (tmp_path / "references.bib").read_text().count("@misc{") == 1
    provenance = json.loads((tmp_path / "provenance.json").read_text())
    assert set(provenance["citations"]) == {"test2020", "alternate"}


def test_literal_citation_tokens_are_not_resolved(tmp_path: Path) -> None:
    report = replace(snapshot(), blocks=(Block("paragraph", (TextSpan("{{cite:missing}} 123.45"),)),))
    result = DocxExporter().export(report, context(tmp_path), {})
    assert result.files
    assert "{{cite:missing}} 123.45" in Document(str(tmp_path / "report.docx")).paragraphs[1].text


def test_latex_options_applied(tmp_path: Path) -> None:
    LatexExporter().export(snapshot(), context(tmp_path), {"english_font": "Example Font", "margin_mm": 30})
    content = (tmp_path / "report.tex").read_text()
    assert r"\setmainfont{Example Font}" in content and "margin=30mm" in content


def test_docx_control_characters_do_not_crash(tmp_path: Path) -> None:
    DocxExporter().export(replace(snapshot(), title="Synthetic\x00 title\x0b"), context(tmp_path), {})
    assert Document(str(tmp_path / "report.docx")).paragraphs[0].text == "Synthetic title"


def test_slides_defaults_and_overflow(tmp_path: Path) -> None:
    from pptx import Presentation

    text = "\n".join(f"Source line {i}" for i in range(23))
    report = replace(
        snapshot(), blocks=(Block("heading", (TextSpan("Results"),)), Block("paragraph", (TextSpan(text),)))
    )
    result = SlidesExporter().export(report, context(tmp_path), {})
    deck = Presentation(str(tmp_path / "report.pptx"))
    titles = [slide.shapes[0].text for slide in deck.slides]
    expected = ["Background", "Objectives", "Method", "Results", "Discussion", "Conclusion"]
    assert [t for t in titles if t in expected] == ["Results"]  # empty sections get no slide (ADR-164)
    assert {d.ref for d in result.diagnostics if d.code == "empty_section"} == set(expected) - {"Results"}
    rendered = []
    for slide in deck.slides:
        assert str(slide.background.fill.fore_color.rgb) == "FFFFFF"
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.name != "Footer":  # the 12 pt footer is the only small text
                lines.extend(shape.text.splitlines())
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        assert run.font.size is not None and run.font.size.pt >= 24
        assert len(lines) <= 10
        rendered.extend(lines)
    assert all(line in rendered for line in text.splitlines())
    assert any(d.code == "slide_overflow" and d.severity == "warning" for d in result.diagnostics)
    sidecar = json.loads((tmp_path / "provenance.json").read_text())
    assert sidecar["diagnostics"] == [d.__dict__ for d in result.diagnostics]


def test_slides_unicode_and_long_title_are_split(tmp_path: Path) -> None:
    from pptx import Presentation

    from researchos_plugin_report_export.slides import LINE_EM, _width

    value = "研究" * 400
    report = replace(snapshot(), title=value, blocks=(Block("paragraph", (TextSpan(value),)),))
    result = SlidesExporter().export(report, context(tmp_path), {})
    deck = Presentation(str(tmp_path / "report.pptx"))
    bodies = []
    for slide in deck.slides:
        lines = slide.shapes[1].text.splitlines()
        assert len(lines) <= 9 and all(_width(line) <= LINE_EM for line in lines)
        bodies.extend(lines)
    assert "".join(bodies).count(value) == 2
    assert any(d.code == "slide_overflow" for d in result.diagnostics)


def test_slides_evidence_and_draft_placeholders(tmp_path: Path) -> None:
    from pptx import Presentation

    ctx = context(tmp_path)
    ctx.mode = "draft"
    report = replace(snapshot(), blocks=(*snapshot().blocks, Block("paragraph", (CiteSpan("missing"),))))
    result = SlidesExporter().export(report, ctx, {})
    deck = Presentation(str(tmp_path / "report.pptx"))
    text = "\n".join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert "[unresolved: missing]" in text and "0.9124" in text and "claim-1" in text
    assert "Synthetic source" in " ".join(text.split())
    assert any(d.code == "unresolved_citation" and d.severity == "warning" for d in result.diagnostics)
    sidecar = json.loads((tmp_path / "provenance.json").read_text())
    assert sidecar["citations"]["test2020"]["evidence"][0]["evidence_id"] == "evidence-1"
    assert sidecar["metrics"]


def test_slides_verified_figure(tmp_path: Path) -> None:
    import base64

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    raw = base64.b64decode((Path(__file__).parents[1] / "fixtures" / "pixel.b64").read_text())
    source = tmp_path / "source.png"
    source.write_bytes(raw)
    ctx = context(tmp_path)
    ctx.artifacts = {"figure": ArtifactRef("figure", "image/png", hashlib.sha256(raw).hexdigest(), "run", source)}
    report = replace(snapshot(), blocks=(Block("figure", artifact_id="figure", caption=(TextSpan("Caption"),)),))
    SlidesExporter().export(report, ctx, {})
    deck = Presentation(str(tmp_path / "report.pptx"))
    pictures = [s for slide in deck.slides for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1 and pictures[0].image.blob == raw
    assert any(s.text == "Caption" for slide in deck.slides for s in slide.shapes if s.has_text_frame)


def test_slides_unsafe_options_and_output(tmp_path: Path) -> None:
    import jsonschema

    with pytest.raises(jsonschema.ValidationError):
        SlidesExporter().export(snapshot(), context(tmp_path), {"font_size": 12})
    outside = tmp_path.parent / "must-not-write.pptx"
    (tmp_path / "report.pptx").symlink_to(outside)
    with pytest.raises(ValueError):
        SlidesExporter().export(snapshot(), context(tmp_path), {})
    assert not outside.exists()


def test_slide_headings_map_to_default_sections_and_words_are_not_split() -> None:
    from researchos_plugin_report_export.resolve import slide_section
    from researchos_plugin_report_export.slides import LINE_EM, _width, slide_lines

    assert [
        slide_section(h)
        for h in ("1. Introduction", "Materials and Methods", "3.2 Results:", "Conclusions", "Data cleaning")
    ] == ["Background", "Method", "Results", "Conclusion", None]
    text = "Heart rate variability decreased in patients compared with healthy controls overall"
    lines = slide_lines(text)
    assert " ".join(lines) == text and all(_width(line) <= LINE_EM for line in lines)
    assert all(word in text.split(" ") for line in lines for word in line.split(" "))  # no word cut in half


def test_slides_keep_list_items_whole_and_bulleted(tmp_path: Path) -> None:
    from pptx import Presentation

    item = "subject-wise split with held-out confirmation subjects and fixed seeds for every model"
    blocks = (
        Block("heading", (TextSpan("Methods"),)),
        *[Block("list", (TextSpan(f"Item {i}: {item}"),)) for i in range(7)],
    )
    SlidesExporter().export(replace(snapshot(), blocks=blocks), context(tmp_path), {})
    pages = [
        s.shapes[1].text.splitlines()
        for s in Presentation(str(tmp_path / "report.pptx")).slides
        if s.shapes[0].text.startswith("Method")
    ]
    assert len(pages) == 2 and all(len(p) <= 9 for p in pages)
    for page in pages:  # every slide starts with a bullet: no item continues from the previous slide
        assert page[0].startswith("• ") and all(x.startswith(("• ", "  ")) for x in page)


def test_slide_lines_fit_the_text_box_when_rendered() -> None:
    """Regression (review #14): measure real glyph widths (Liberation Sans = Arial metrics) at 24 pt."""
    from PIL import ImageFont

    from researchos_plugin_report_export.slides import slide_lines

    font_path = Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf")
    if not font_path.is_file():
        pytest.skip("Liberation Sans not installed")
    font = ImageFont.truetype(str(font_path), 1000)
    usable_pt = (12.3 - 0.2) * 72
    for text in ("W" * 60, " ".join(["WMWM"] * 20), "mmmm " * 15, "研究" * 40, "\u2022 " + "Mixed WIDE text ü " * 6):
        for line in slide_lines(text):
            assert font.getlength(line) / 1000 * 24 <= usable_pt, (text[:12], line)


def test_slides_title_footer_and_native_tables(tmp_path: Path) -> None:
    """slides_export step 2 (ADR-160): title slide, numbered footers, native tables paged by rows, numbers
    right-aligned, every shape inside the slide."""
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu

    rows = tuple(((TextSpan("metric"),), (TextSpan("split"),), (TextSpan("value"),)) for _ in range(1)) + tuple(
        ((TextSpan(f"metric_{i}"),), (TextSpan("search_val"),), (TextSpan(f"0.{i:03d}"),)) for i in range(40)
    )
    report = replace(
        snapshot(), kind="report", blocks=(Block("heading", (TextSpan("Results"),)), Block("table", rows=rows))
    )
    result = SlidesExporter().export(report, context(tmp_path), {})
    deck = Presentation(str(tmp_path / "report.pptx"))
    assert deck.slides[0].shapes[0].text == "Report" and "Synthetic" in deck.slides[0].shapes[1].text
    tables = [s.table for slide in deck.slides for s in slide.shapes if s.has_table]
    assert len(tables) >= 3 and all(t.cell(0, 0).text == "metric" for t in tables)  # header repeats per page
    assert sum(len(t.rows) - 1 for t in tables) == 40
    cell = tables[0].cell(1, 2)
    assert (
        cell.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
        and tables[0].cell(0, 0).text_frame.paragraphs[0].runs[0].font.bold
    )
    n = len(deck.slides)
    for i, slide in enumerate(deck.slides, start=1):
        footers = [s for s in slide.shapes if s.name == "Footer"]
        assert footers[0].text.endswith(f"{i} / {n}") if i > 1 else not footers
        for s in slide.shapes:
            assert s.left >= 0 and s.top >= 0 and s.left + s.width <= deck.slide_width + Emu(1)
            assert s.top + s.height <= deck.slide_height + Emu(1), (i, s.name)
    assert any(d.code == "slide_overflow" for d in result.diagnostics)


def test_paper_cards_export_as_slides(tmp_path: Path) -> None:
    result = SlidesExporter().export(replace(snapshot(), kind="paper_card"), context(tmp_path), {})
    assert result.files and SlidesExporter.spec.doc_kinds[-1] == "paper_card"


def test_tall_rows_continue_and_tables_end_above_the_footer(tmp_path: Path) -> None:
    """Regression (review #62): one long description cell ended at 19.33 in on a 7.5 in slide."""
    from pptx import Presentation
    from pptx.util import Inches

    long = " ".join(f"word{i}" for i in range(600))
    rows = (
        ((TextSpan("Item"),), (TextSpan("Description"),)),
        ((TextSpan("a"),), (TextSpan(long),)),
        ((TextSpan("b"),), (TextSpan("short"),)),
    )
    report = replace(snapshot(), blocks=(Block("heading", (TextSpan("Results"),)), Block("table", rows=rows)))
    result = SlidesExporter().export(report, context(tmp_path), {})
    deck = Presentation(str(tmp_path / "report.pptx"))
    tables = [s for slide in deck.slides for s in slide.shapes if s.has_table]
    assert len(tables) >= 3 and any(d.code == "slide_overflow" for d in result.diagnostics)
    for t in tables:
        assert t.top + sum(r.height for r in t.table.rows) <= Inches(6.9) + 1
    text = " ".join(t.table.cell(r, 1).text for t in tables for r in range(1, len(t.table.rows)))
    assert all(f"word{i}" in text.split() for i in (0, 299, 599)) and "short" in text  # nothing lost
